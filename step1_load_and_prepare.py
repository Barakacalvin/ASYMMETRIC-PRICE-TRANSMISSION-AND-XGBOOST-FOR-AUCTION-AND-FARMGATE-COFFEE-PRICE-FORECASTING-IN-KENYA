"""
step1_load_and_prepare.py
-------------------------
Load the master workbook, harmonise units to USD/kg, align all monthly series
to the Kenyan coffee year (October = month 1, September = month 12), build the
26 x 12 = 312 monthly panel and the 26-year annual panel.

Implements Chapter 3 v3 sections 3.3.1, 3.3.2 and 3.3.3.

Outputs
-------
Codes/_outputs/data/monthly_clean.csv   (312 rows: monthly USD/kg series)
Codes/_outputs/data/annual_clean.csv    (26 rows : annual USD/kg series)
Codes/_outputs/data/data_audit.json     (provenance + recess-month counts n_y)
"""
from __future__ import annotations
import json
import warnings
import numpy as np
import pandas as pd
import openpyxl
from pathlib import Path

import config as C
from logging_setup import get_logger
log = get_logger(__name__)

warnings.filterwarnings("ignore")

# ============================================================ helpers
def _read_indicator_wide(ws, indicator: str) -> pd.DataFrame:
    """Read one indicator from the wide ICIP sheet into a long monthly frame."""
    rows = []
    for r in range(2, ws.max_row + 1):
        ind = ws.cell(r, 1).value
        if ind != indicator:
            continue
        year = ws.cell(r, 2).value
        if year is None:
            continue
        for c, m in enumerate(range(1, 13), start=3):
            v = ws.cell(r, c).value
            rows.append({"calendar_year": int(year),
                         "calendar_month": m,
                         "value": v})
    df = pd.DataFrame(rows)
    return df


def _to_coffee_year_index(calendar_year: int, calendar_month: int):
    """
    Map (calendar_year, calendar_month) to ('YYYY/YYYY', coffee_year_month).
    Coffee year y/y+1 spans Oct of y through Sep of y+1.
    October = month 1 ... September = month 12.
    """
    if calendar_month >= 10:
        cy_start = calendar_year
        cy_month = calendar_month - 9       # Oct -> 1
    else:
        cy_start = calendar_year - 1
        cy_month = calendar_month + 3       # Jan -> 4 ... Sep -> 12
    return f"{cy_start}/{cy_start + 1}", cy_month


def _coffee_year_to_int(cy: str) -> int:
    """'1998/1999' -> 1998   (the start year)."""
    return int(cy.split("/")[0])


def _sample_filter(df: pd.DataFrame) -> pd.DataFrame:
    start = _coffee_year_to_int(C.SAMPLE_START_COFFEE_YEAR)
    end   = _coffee_year_to_int(C.SAMPLE_END_COFFEE_YEAR)
    cy_int = df["coffee_year"].apply(_coffee_year_to_int)
    mask = (cy_int >= start) & (cy_int <= end)
    return df.loc[mask].reset_index(drop=True)


# ============================================================ load
def load_workbook():
    log.info(f"Reading: {C.WORKBOOK}")
    wb = openpyxl.load_workbook(C.WORKBOOK, data_only=True)
    return wb


def build_monthly_panel(wb) -> pd.DataFrame:
    """Build the 312-row monthly panel in USD/kg, coffee-year indexed."""
    icip = wb["ICIP_NY-Futures_Auct_Columbian"]
    fxws = wb["Exchange_Rate_Data"]

    # --- international references (cents/lb -> USD/kg) ---------------
    series_cfg = [
        ("New York Futures",          "ny_usdkg",   C.CENTS_PER_LB_TO_USD_PER_KG),
        ("ICO Composite Indicator",   "ico_usdkg",  C.CENTS_PER_LB_TO_USD_PER_KG),
        ("Columbian Milds",           "col_usdkg",  C.CENTS_PER_LB_TO_USD_PER_KG),
        ("Auction Monthly Averages",  "auc_usdkg",  C.USD_PER_50KG_TO_USD_PER_KG),
    ]
    monthly = None
    recess_flags = None
    for indicator, colname, factor in series_cfg:
        d = _read_indicator_wide(icip, indicator)
        # Recess strings in auction: convert to NaN and remember the flag
        if indicator == "Auction Monthly Averages":
            d["recess_flag"] = d["value"].apply(
                lambda v: True if isinstance(v, str) and v.strip().lower() == "recess" else False
            )
            d["value"] = d["value"].apply(
                lambda v: np.nan if isinstance(v, str) else v
            )
        d["value"] = pd.to_numeric(d["value"], errors="coerce") * factor
        d = d.rename(columns={"value": colname})
        keep = ["calendar_year", "calendar_month", colname]
        if "recess_flag" in d.columns:
            keep.append("recess_flag")
        d = d[keep]
        if monthly is None:
            monthly = d
        else:
            monthly = monthly.merge(d, on=["calendar_year", "calendar_month"], how="outer")

    # --- exchange rate (USD/KES already a rate, not a price) ---------
    fx_rows = []
    for r in range(2, fxws.max_row + 1):
        cy_str = fxws.cell(r, 5).value
        cym    = fxws.cell(r, 6).value
        cyear  = fxws.cell(r, 3).value
        cmonth = fxws.cell(r, 4).value
        usdkes = fxws.cell(r, 7).value
        if usdkes is None:
            continue
        fx_rows.append({
            "calendar_year"  : int(cyear),
            "calendar_month" : int(cmonth),
            "fx_usdkes"      : float(usdkes),
        })
    fx = pd.DataFrame(fx_rows)
    monthly = monthly.merge(fx, on=["calendar_year", "calendar_month"], how="left")

    # --- coffee-year indexing ----------------------------------------
    monthly[["coffee_year", "coffee_year_month"]] = monthly.apply(
        lambda r: pd.Series(_to_coffee_year_index(int(r["calendar_year"]),
                                                   int(r["calendar_month"]))),
        axis=1,
    )

    # ordering and filter
    monthly["recess_flag"] = monthly.get("recess_flag", False).fillna(False).astype(bool)
    monthly = monthly[[
        "coffee_year", "coffee_year_month",
        "calendar_year", "calendar_month",
        "ny_usdkg", "ico_usdkg", "col_usdkg",
        "auc_usdkg", "recess_flag",
        "fx_usdkes",
    ]].sort_values(["coffee_year", "coffee_year_month"]).reset_index(drop=True)

    monthly = _sample_filter(monthly)
    return monthly


def build_annual_panel(wb, monthly: pd.DataFrame) -> pd.DataFrame:
    """
    Build the 26-row annual panel:
    - farmgate clean USD/kg (from Producer Prices sheet, AFA Yearbook)
    - cherry farmgate (KES/kg, 5 obs only — descriptive use)
    - within-coffee-year mean of monthly auction (USD/kg), recess months excluded
    - within-coffee-year means of NY/ICO/Colombian (USD/kg)
    - within-coffee-year mean USD/KES rate
    - n_auction_months_y for transparency
    """
    pp = wb["Producer Prices"]
    rows = []
    for r in range(2, pp.max_row + 1):
        cy = pp.cell(r, 1).value
        if cy is None or "/" not in str(cy):
            continue
        rows.append({
            "coffee_year"      : str(cy),
            "fg_usdkg_clean"   : pp.cell(r, 3).value,
            "fg_keskg_clean"   : pp.cell(r, 4).value,
            "fx_year_avg"      : pp.cell(r, 2).value,
            "cherry_keskg"     : pp.cell(r, 6).value,
        })
    annual = pd.DataFrame(rows)
    annual = _sample_filter(annual)
    annual["fg_usdkg_clean"] = pd.to_numeric(annual["fg_usdkg_clean"], errors="coerce")

    # within-CY auction mean, recess excluded
    nrecess = (monthly.assign(
        is_recess=monthly["recess_flag"] | monthly["auc_usdkg"].isna()
    ).groupby("coffee_year")["is_recess"].sum().reset_index()
    .rename(columns={"is_recess": "n_recess_months_y"}))

    grp = (monthly.dropna(subset=["auc_usdkg"])
                  .groupby("coffee_year")
                  .agg(auc_usdkg_mean=("auc_usdkg", "mean"),
                       n_auction_months_y=("auc_usdkg", "size"))
                  .reset_index())
    annual = annual.merge(grp, on="coffee_year", how="left")
    annual = annual.merge(nrecess, on="coffee_year", how="left")

    # within-CY means for international references and FX
    for src, name in [("ny_usdkg", "ny_usdkg_mean"),
                      ("ico_usdkg", "ico_usdkg_mean"),
                      ("col_usdkg", "col_usdkg_mean"),
                      ("fx_usdkes", "fx_usdkes_mean")]:
        g = monthly.dropna(subset=[src]).groupby("coffee_year")[src].mean().reset_index()
        g = g.rename(columns={src: name})
        annual = annual.merge(g, on="coffee_year", how="left")

    annual = annual.sort_values("coffee_year").reset_index(drop=True)
    return annual


# ============================================================ main
def main():
    wb = load_workbook()
    monthly = build_monthly_panel(wb)
    annual  = build_annual_panel(wb, monthly)

    log.info(f"Monthly panel: {len(monthly):4d} rows   "
          f"(expected {C.N_MONTHS_EXPECTED}, range "
          f"{monthly['coffee_year'].iloc[0]} .. {monthly['coffee_year'].iloc[-1]})")
    log.info(f"Annual panel : {len(annual):4d} rows   "
          f"(expected {C.N_COFFEE_YEARS_EXPECTED})")

    # provenance / audit
    audit = {
        "workbook"            : str(C.WORKBOOK),
        "sample_start"        : C.SAMPLE_START_COFFEE_YEAR,
        "sample_end"          : C.SAMPLE_END_COFFEE_YEAR,
        "n_monthly_rows"      : int(len(monthly)),
        "n_annual_rows"       : int(len(annual)),
        "n_auction_months_y"  : annual.set_index("coffee_year")["n_auction_months_y"].to_dict(),
        "n_recess_months_y"   : annual.set_index("coffee_year")["n_recess_months_y"].to_dict(),
        "unit_convention"     : "USD/kg (clean coffee equivalent) for all price series",
        "fx_unit"             : "KES per USD",
        "international_primary"      : "New York 'C' Futures",
        "international_robustness"   : ["ICO Composite Indicator", "Columbian Milds"],
        "farmgate_target"            : "AFA Yearbook clean-equivalent (USD/kg)",
        "cherry_role"                : "5-obs descriptive comparison only (Sec 3.7.5)",
    }

    monthly.to_csv(C.DATA_OUT / "monthly_clean.csv", index=False)
    annual.to_csv (C.DATA_OUT / "annual_clean.csv",  index=False)
    with open(C.DATA_OUT / "data_audit.json", "w") as f:
        json.dump(audit, f, indent=2, default=str)

    log.info(f"\nWrote {C.DATA_OUT/'monthly_clean.csv'}")
    log.info(f"Wrote {C.DATA_OUT/'annual_clean.csv'}")
    log.info(f"Wrote {C.DATA_OUT/'data_audit.json'}")


if __name__ == "__main__":
    main()
